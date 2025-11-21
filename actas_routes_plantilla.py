# actas_routes.py — generación de ACTAS en DOCX (con plantilla Mediazion)
from fastapi import APIRouter, HTTPException, Request
from pydantic import BaseModel
from typing import Optional
import os
from datetime import datetime

from docx import Document

actas_router = APIRouter()

# Directorio donde se guardan las actas generadas
ACTAS_DIR = os.path.join("uploads", "actas")
os.makedirs(ACTAS_DIR, exist_ok=True)

# Ruta de la plantilla base (DOCX) con el diseño azul + logo
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(BASE_DIR, "templates", "ActaBaseMediazion.docx")


class ActaIn(BaseModel):
    case_no: str
    date_iso: str
    mediator_alias: str
    parties: str
    summary: str
    agreements: Optional[str] = None
    confidentiality: Optional[bool] = True
    logo_url: Optional[str] = "https://mediazion.eu/logo.png"  # se mantiene por compatibilidad, pero la plantilla ya incluye logo


def _replace_placeholders(doc: Document, mapping: dict):
    """Reemplaza marcadores simples {{KEY}} en todos los párrafos del documento."""
    for p in doc.paragraphs:
        if not p.text:
            continue
        for key, value in mapping.items():
            if key in p.text:
                inline_text = "".join(run.text for run in p.runs)
                inline_text = inline_text.replace(key, value)
                for i in range(len(p.runs) - 1, -1, -1):
                    p.runs[i].text = ""
                if p.runs:
                    p.runs[0].text = inline_text


@actas_router.post("/actas/render_docx")
def render_docx(body: ActaIn, request: Request):
    """Genera un ACTA en DOCX usando la plantilla base de Mediazion."""
    try:
        if not os.path.exists(TEMPLATE_PATH):
            raise HTTPException(500, f"No se encontró la plantilla de actas en: {TEMPLATE_PATH}")

        doc = Document(TEMPLATE_PATH)

        conf_text = ""
        if body.confidentiality:
            conf_text = (
                "Las partes se comprometen a mantener la confidencialidad del proceso de mediación "
                "y de la información intercambiada, salvo obligación legal o acuerdo expreso en contrario."
            )

        mapping = {
            "{{CASE_NO}}": body.case_no,
            "{{DATE_ISO}}": body.date_iso,
            "{{MEDIATOR}}": body.mediator_alias,
            "{{PARTIES}}": body.parties,
            "{{SUMMARY}}": body.summary,
            "{{AGREEMENTS}}": body.agreements or "Sin acuerdos adicionales registrados.",
            "{{CONF_TEXT}}": conf_text,
        }

        _replace_placeholders(doc, mapping)

        ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        safe_case = body.case_no.replace(" ", "_")
        filename = f"Acta_{safe_case}_{ts}.docx"
        path = os.path.join(ACTAS_DIR, filename)
        doc.save(path)

        base_url = str(request.base_url).rstrip("/")
        return {"ok": True, "url": f"{base_url}/uploads/actas/{filename}"}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Error generando el acta: {str(e)}")
